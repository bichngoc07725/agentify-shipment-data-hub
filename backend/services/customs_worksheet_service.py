"""Phiếu nhập liệu tờ khai hải quan (Bước 3) — file .docx để gõ sang ECUS/VNACCS.

Agentify KHÔNG nối vào ECUS: đó là phần mềm ngoài, có quy trình chữ ký số riêng.
Việc ở đây khiêm tốn hơn nhưng đúng chỗ đau: gom mọi thứ hệ thống đã biết về lô
hàng, xếp theo ĐÚNG THỨ TỰ các ô trên tờ khai, để nhân viên nhìn một tờ mà gõ
thay vì mở song song năm cái email và ba file PDF.

Nguyên tắc chi phối toàn bộ module: ô nào không có dữ liệu thì ghi rõ "KHÔNG CÓ
TRONG AGENTIFY — cần lấy từ <nguồn>", không để trống. Một ô trống trên tờ giấy
in ra trông y hệt một ô đã kiểm và đúng là rỗng; nhầm hai thứ đó ở khâu khai hải
quan dẫn tới tờ khai sai, và sửa tờ khai sau khi truyền là chuyện tốn tiền.
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from db.models import (
    Booking,
    Container,
    ContainerFact,
    CustomsDeclaration,
    Quote,
)

MISSING_PREFIX = "KHÔNG CÓ TRONG AGENTIFY"


@dataclass(frozen=True)
class WorksheetField:
    """Một ô trên tờ khai."""

    label: str
    value: str | None
    # Chỗ nhân viên phải đi lấy khi ô này trống. Nói ra được thì tiết kiệm đúng
    # cái vòng "không biết thiếu ở đâu" mà tờ khai hay mắc.
    source_hint: str = ""

    @property
    def display(self) -> str:
        if self.value:
            return self.value
        return f"{MISSING_PREFIX}{f' — cần lấy từ {self.source_hint}' if self.source_hint else ''}"

    @property
    def is_missing(self) -> bool:
        return not self.value


@dataclass(frozen=True)
class WorksheetSection:
    title: str
    fields: tuple[WorksheetField, ...]


async def _facts_for(db: AsyncSession, container_id) -> dict[str, str]:
    """Giá trị mới nhất của từng field trong hồ sơ container."""
    result = await db.execute(
        select(
            ContainerFact.field_name,
            ContainerFact.normalized_value,
            ContainerFact.field_value,
        )
        .where(ContainerFact.container_id == container_id)
        .order_by(
            ContainerFact.source_sent_at.desc().nullslast(),
            ContainerFact.created_at.desc(),
        )
    )
    facts: dict[str, str] = {}
    for field_name, normalized, raw in result.all():
        facts.setdefault(field_name, normalized or raw)
    return facts


def _fmt(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def build_sections(
    container: Container,
    quote: Quote | None,
    booking: Booking | None,
    declaration: CustomsDeclaration | None,
    facts: dict[str, str],
) -> list[WorksheetSection]:
    """Các ô của tờ khai, xếp theo thứ tự người khai đi trên màn hình ECUS."""

    def q(attr: str):
        return _fmt(getattr(quote, attr, None)) if quote else None

    def b(attr: str):
        return _fmt(getattr(booking, attr, None)) if booking else None

    return [
        WorksheetSection(
            "A. Thông tin chung tờ khai",
            (
                WorksheetField(
                    "Số tờ khai",
                    _fmt(declaration.declaration_no) if declaration else None,
                    "ECUS cấp sau khi truyền",
                ),
                WorksheetField(
                    "Loại hình",
                    _fmt(declaration.declaration_type.value) if declaration else None,
                    "chọn trên ECUS",
                ),
                WorksheetField("Số vận đơn (B/L)", _fmt(facts.get("bl_no")), "email hãng tàu / vận đơn"),
                WorksheetField("Số booking", b("booking_no") or _fmt(facts.get("booking_no")), "xác nhận đặt chỗ"),
            ),
        ),
        WorksheetSection(
            "B. Bên liên quan",
            (
                WorksheetField(
                    "Người xuất khẩu — tên",
                    _fmt(facts.get("shipper")) or q("customer_name") or _fmt(facts.get("customer_name")),
                    "báo giá / hợp đồng",
                ),
                WorksheetField("Người xuất khẩu — mã số thuế", _fmt(facts.get("shipper_tax_code")), "hồ sơ khách hàng"),
                WorksheetField("Người xuất khẩu — địa chỉ", _fmt(facts.get("shipper_address")), "hồ sơ khách hàng"),
                WorksheetField(
                    "Người nhập khẩu — tên, địa chỉ, nước",
                    _fmt(facts.get("consignee")),
                    "hoá đơn thương mại",
                ),
            ),
        ),
        WorksheetSection(
            "C. Vận chuyển",
            (
                WorksheetField("Tên tàu", b("vessel") or _fmt(facts.get("vessel")), "xác nhận đặt chỗ"),
                WorksheetField("Số chuyến", b("voyage") or _fmt(facts.get("voyage")), "xác nhận đặt chỗ"),
                WorksheetField("Hãng tàu", b("carrier"), "xác nhận đặt chỗ"),
                WorksheetField("Cảng xếp hàng (POL)", b("pol") or q("pol") or _fmt(facts.get("pol")), "báo giá"),
                WorksheetField("Cảng dỡ hàng (POD)", b("pod") or q("pod") or _fmt(facts.get("pod")), "báo giá"),
                WorksheetField("Ngày khởi hành dự kiến (ETD)", b("etd") or _fmt(facts.get("etd")), "xác nhận đặt chỗ"),
                WorksheetField("Ngày đến dự kiến (ETA)", b("eta") or _fmt(facts.get("eta")), "xác nhận đặt chỗ"),
                WorksheetField("Số container", _fmt(container.container_no), ""),
                WorksheetField("Số seal", _fmt(facts.get("seal_no")), "phiếu EIR / ảnh hiện trường"),
                WorksheetField("Chi cục hải quan", _fmt(facts.get("customs_office")), "thông báo hải quan"),
            ),
        ),
        WorksheetSection(
            "D. Hoá đơn & trị giá",
            (
                WorksheetField("Số hoá đơn thương mại", _fmt(facts.get("invoice_no")), "hoá đơn thương mại"),
                WorksheetField("Ngày hoá đơn", _fmt(facts.get("doc_date")), "hoá đơn thương mại"),
                WorksheetField("Điều kiện giao hàng (Incoterm)", q("incoterm"), "báo giá"),
                WorksheetField("Đồng tiền thanh toán", q("currency"), "báo giá"),
                WorksheetField(
                    "Phương thức thanh toán",
                    q("payment_term") or _fmt(facts.get("payment_term")),
                    "báo giá / hợp đồng",
                ),
                # Cố ý để trống thay vì lấy tổng báo giá: tổng báo giá là tiền
                # CƯỚC bán cho khách, không phải trị giá lô hàng khai hải quan.
                # Điền nhầm con số này vào tờ khai là khai sai trị giá.
                WorksheetField("Tổng trị giá hoá đơn", None, "hoá đơn thương mại (KHÔNG dùng tổng báo giá)"),
            ),
        ),
        WorksheetSection(
            "E. Hàng hoá",
            (
                WorksheetField("Mô tả hàng hoá", q("commodity"), "báo giá"),
                WorksheetField(
                    "Mã HS",
                    (_fmt(declaration.hs_code) if declaration else None) or _fmt(facts.get("hs_code")),
                    "hoá đơn thương mại / tra biểu thuế",
                ),
                WorksheetField("Xuất xứ", None, "C/O hoặc hoá đơn"),
                WorksheetField(
                    "Loại & số lượng container",
                    " x ".join(filter(None, [b("container_qty") or q("container_qty"), b("container_type") or q("container_type")]))
                    or None,
                    "báo giá / đặt chỗ",
                ),
                WorksheetField(
                    "Trọng lượng gộp (KGS)",
                    _fmt(facts.get("gross_weight_kg")) or q("gross_weight_kg"),
                    "packing list",
                ),
                WorksheetField("Số kiện / loại kiện", _fmt(facts.get("packages")), "packing list"),
                WorksheetField("Số khối (CBM)", _fmt(facts.get("volume_cbm")), "packing list"),
                WorksheetField("Đơn giá / trị giá dòng hàng", None, "hoá đơn thương mại"),
            ),
        ),
    ]


async def collect_worksheet(
    db: AsyncSession, container: Container
) -> list[WorksheetSection]:
    facts = await _facts_for(db, container.id)

    booking = (
        await db.execute(
            select(Booking)
            .options(selectinload(Booking.quote))
            .where(Booking.container_id == container.id)
            .order_by(Booking.created_at.desc())
            .limit(1)
        )
    ).scalar_one_or_none()

    quote = booking.quote if booking and booking.quote else None
    if quote is None:
        quote = (
            await db.execute(
                select(Quote)
                .where(Quote.container_id == container.id)
                .order_by(Quote.created_at.desc())
                .limit(1)
            )
        ).scalar_one_or_none()

    declaration = (
        await db.execute(
            select(CustomsDeclaration)
            .where(CustomsDeclaration.container_id == container.id)
            .order_by(CustomsDeclaration.created_at.desc())
            .limit(1)
        )
    ).scalar_one_or_none()

    return build_sections(container, quote, booking, declaration, facts)


def render_docx(container_no: str, sections: list[WorksheetSection]) -> bytes:
    """Dựng file .docx: mỗi mục một bảng hai cột, ô thiếu tô đỏ."""
    from docx import Document
    from docx.shared import Pt, RGBColor

    document = Document()
    document.add_heading(f"PHIẾU NHẬP LIỆU TỜ KHAI HẢI QUAN — {container_no}", level=1)

    intro = document.add_paragraph()
    intro.add_run(
        "Phiếu này để gõ sang ECUS/VNACCS. Agentify không nối trực tiếp vào ECUS. "
    )
    warn = intro.add_run(
        "Ô ghi 'KHÔNG CÓ TRONG AGENTIFY' là ô hệ thống chưa có dữ liệu — "
        "phải tự tra theo nguồn ghi kèm, KHÔNG được bỏ trống trên tờ khai."
    )
    warn.bold = True

    missing_total = 0
    for section in sections:
        document.add_heading(section.title, level=2)
        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for field in section.fields:
            row = table.add_row().cells
            label_run = row[0].paragraphs[0].add_run(field.label)
            label_run.bold = True
            value_run = row[1].paragraphs[0].add_run(field.display)
            if field.is_missing:
                missing_total += 1
                value_run.font.color.rgb = RGBColor(0xC0, 0x30, 0x30)
                value_run.italic = True
            row[0].width = Pt(200)

    footer = document.add_paragraph()
    footer.add_run(
        f"Tổng cộng {missing_total} ô cần tự bổ sung. "
        "Kiểm lại toàn bộ trước khi truyền tờ khai."
    ).italic = True

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
